from docx import Document

from job_agent.documents import extract_resume_text, write_documents


def test_extract_docx_and_generate_word_documents(tmp_path, monkeypatch):
    source = tmp_path / "master.docx"
    document = Document()
    document.add_paragraph("MSc Computational Engineering")
    document.save(source)
    assert "Computational Engineering" in extract_resume_text(source)

    monkeypatch.setattr("job_agent.documents.DATA_DIR", tmp_path / "private")
    payload = {
        "resume_title": "Data Analyst",
        "summary": "Analyst with Python and SQL experience.",
        "skills": ["Python", "SQL"],
        "experience": [
            {
                "role": "Data Engineering Intern",
                "company": "Razorpay",
                "dates": "",
                "bullets": ["Worked with data pipelines."],
            }
        ],
        "education": ["MSc Computational Engineering"],
        "cover_letter": "Dear hiring team, I am applying for this role.",
    }
    resume, cover = write_documents(3, "Example GmbH", payload)
    assert resume.is_file() and cover.is_file()
    assert "Data Analyst" in "\n".join(p.text for p in Document(resume).paragraphs)
    assert "Dear hiring team" in Document(cover).paragraphs[0].text

