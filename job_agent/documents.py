from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
import uuid
from typing import Any

from docx import Document
from pypdf import PdfReader

from .db import DATA_DIR

MAX_RESUME_BYTES = 15 * 1024 * 1024


def extract_resume_text(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError("The selected resume file does not exist.")
    if path.stat().st_size > MAX_RESUME_BYTES:
        raise ValueError("Resume files must be 15 MB or smaller.")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    elif suffix == ".docx":
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(cell.text for row in table.rows for cell in row.cells)
        text = "\n".join(paragraphs)
    else:
        raise ValueError("Upload a PDF or DOCX resume.")
    if not text.strip():
        raise ValueError("No readable text was found in the resume.")
    return text.strip()


def save_uploaded_resume(name: str, content: bytes) -> Path:
    suffix = Path(name).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise ValueError("Upload a PDF or DOCX resume.")
    if len(content) > MAX_RESUME_BYTES:
        raise ValueError("Resume files must be 15 MB or smaller.")
    directory = DATA_DIR / "resumes"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"master-{uuid.uuid4().hex}{suffix}"
    path.write_bytes(content)
    return path


def _safe_name(value: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_-]+", "-", value).strip("-")
    return name[:60] or "application"


def write_documents(job_id: int, company: str, payload: dict[str, Any]) -> tuple[Path, Path]:
    directory = DATA_DIR / "generated"
    directory.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    stem = f"job-{job_id}-{_safe_name(company)}-{stamp}"

    resume = Document()
    resume.add_heading(str(payload["resume_title"]), level=0)
    resume.add_heading("Professional Summary", level=1)
    resume.add_paragraph(str(payload["summary"]))
    resume.add_heading("Skills", level=1)
    resume.add_paragraph(", ".join(str(skill) for skill in payload["skills"]))
    resume.add_heading("Experience", level=1)
    for entry in payload["experience"]:
        resume.add_heading(
            " | ".join(
                str(entry.get(field, "")).strip()
                for field in ("role", "company", "dates")
                if entry.get(field)
            ),
            level=2,
        )
        for bullet in entry.get("bullets", []):
            resume.add_paragraph(str(bullet), style="List Bullet")
    resume.add_heading("Education", level=1)
    for item in payload["education"]:
        resume.add_paragraph(str(item), style="List Bullet")
    resume_path = directory / f"{stem}-resume.docx"
    resume.save(resume_path)

    letter = Document()
    letter.add_paragraph(str(payload["cover_letter"]))
    cover_path = directory / f"{stem}-cover-letter.docx"
    letter.save(cover_path)
    return resume_path, cover_path

